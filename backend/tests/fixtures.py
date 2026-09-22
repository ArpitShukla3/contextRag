"""Fixtures generating in-memory documents for parsing tests."""

import io

import pymupdf
import pytest
from docx import Document as DocxDocument


@pytest.fixture()
def pdf_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF")
    page.insert_text((72, 100), "Second PDF line")
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture()
def multi_page_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    for text in ("Page one content", "Page two content"):
        page = doc.new_page()
        page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture()
def txt_bytes() -> bytes:
    return "Plain text content.\nSecond line.".encode("utf-8")


@pytest.fixture()
def markdown_bytes() -> bytes:
    return b"# Heading\n\nSome **bold** markdown text."


@pytest.fixture()
def docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Hello DOCX")
    document.add_paragraph("Second paragraph")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()