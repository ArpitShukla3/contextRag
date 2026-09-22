"""DOCX document parser (python-docx)."""

from app.exceptions import ParsingError
from app.services.parsing.base import DocumentParser, ParsedDocument, ParsedPage

try:
    import io

    from docx import Document as DocxDocument
except ImportError as exc:  # pragma: no cover - dependency missing
    raise ImportError("python-docx is required for DOCX parsing.") from exc


class DocxParser(DocumentParser):
    """Parser for Microsoft Word ``.docx`` files.

    Page numbers are not available from the OOXML flow layout, so pages are
    reported as a single continuous region.
    """

    extensions = frozenset({".docx"})
    content_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
    )

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            document = DocxDocument(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
        except Exception as exc:
            raise ParsingError(
                "Could not open DOCX: the file is corrupt or not a Word document."
            ) from exc

        text = "\n".join(paragraphs)
        return ParsedDocument(
            pages=[ParsedPage(text=text)],
            metadata={"format": "docx"},
        )